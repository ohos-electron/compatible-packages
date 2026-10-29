#!/usr/bin/env python3
"""
从 Electron 兼容性分析结果 xlsx 中读取「状态 = 通过」的 npm 包，生成：
- data/*.md        Markdown 表格（GitHub 在线浏览）
- data/*.jsonl     JSONL 格式（命令行 grep/jq 查询，搜索页数据源）
- data/index.json  索引文件（总数、更新日期、各文件条目数）
- README.md        仓库首页

用法：python3 gen_packages.py <xlsx文件路径> [--repo owner/name] [--electron 34]
"""

import argparse
import glob
import json
import os
import sys
from collections import defaultdict
from datetime import date
from openpyxl import load_workbook

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(SCRIPT_DIR, "data")
README_FILE = os.path.join(SCRIPT_DIR, "README.md")

SHEET_NAME = "包结论汇总"
COL_STATUS = "状态"
COL_NAME = "包名"
COL_VERSION = "版本"
PASS_STATUS = "通过"
LETTERS = list("abcdefghijklmnopqrstuvwxyz") + ["0-9"]


def get_sort_key(name):
    if name.startswith("@"):
        parts = name.split("/", 1)
        if len(parts) == 2:
            return (parts[0][1:].lower() + "/" + parts[1].lower(),)
        return (name.lower(),)
    return (name.lower(),)


def get_first_letter(name):
    if name.startswith("@"):
        after_at = name[1:]
        if after_at:
            c = after_at[0].lower()
            return c if "a" <= c <= "z" else "0-9"
    c = name[0].lower() if name else "_"
    return c if "a" <= c <= "z" else "0-9"


def cell_str(value):
    return str(value).strip() if value is not None else ""


def read_packages(xlsx_path):
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    rows = ws.iter_rows(values_only=True)

    header = [cell_str(h) for h in next(rows)]
    missing = [c for c in (COL_STATUS, COL_NAME, COL_VERSION) if c not in header]
    if missing:
        sys.exit(f"表头缺少列：{', '.join(missing)}（sheet：{ws.title}）")
    i_status = header.index(COL_STATUS)
    i_name = header.index(COL_NAME)
    i_version = header.index(COL_VERSION)

    packages = {}
    skipped = 0
    for row in rows:
        if not row or len(row) <= max(i_status, i_name, i_version):
            continue
        name = cell_str(row[i_name])
        if not name:
            continue
        if cell_str(row[i_status]) != PASS_STATUS:
            skipped += 1
            continue
        packages[name] = {"name": name, "version": cell_str(row[i_version]) or "latest"}
    wb.close()

    print(f"Sheet「{ws.title}」：通过 {len(packages)} 个，跳过 {skipped} 个（失败/受阻/有条件通过）")
    return sorted(packages.values(), key=lambda p: get_sort_key(p["name"]))


def write_data(packages, electron_version, repo):
    groups = defaultdict(list)
    for p in packages:
        groups[get_first_letter(p["name"])].append(p)

    os.makedirs(DATA_DIR, exist_ok=True)
    for old in glob.glob(os.path.join(DATA_DIR, "*.jsonl")) + glob.glob(os.path.join(DATA_DIR, "*.md")):
        os.remove(old)

    file_counts = {}
    for letter in LETTERS:
        entries = groups.get(letter)
        if not entries:
            continue
        with open(os.path.join(DATA_DIR, f"{letter}.md"), "w", encoding="utf-8") as f:
            f.write(f"# 字母 {letter.upper()} 开头的兼容包（{len(entries)} 个）\n\n")
            f.write("> [← 返回首页](../README.md)\n\n")
            f.write("| 包名 | 版本 |\n")
            f.write("|------|------|\n")
            for p in entries:
                f.write(f"| `{p['name']}` | {p['version']} |\n")

        with open(os.path.join(DATA_DIR, f"{letter}.jsonl"), "w", encoding="utf-8") as f:
            for p in entries:
                f.write(json.dumps(p, ensure_ascii=False) + "\n")

        file_counts[f"{letter}.jsonl"] = len(entries)
        print(f"  {letter}: {len(entries)}")

    index_data = {
        "total": len(packages),
        "last_updated": str(date.today()),
        "electron_version": electron_version,
        "repo": repo,
        "files": file_counts,
    }
    with open(os.path.join(DATA_DIR, "index.json"), "w", encoding="utf-8") as f:
        json.dump(index_data, f, ensure_ascii=False, indent=2)
        f.write("\n")
    return index_data


def write_readme(index_data):
    repo = index_data["repo"]
    owner, name = repo.split("/", 1)
    site_url = f"https://{owner.lower()}.github.io/{name}/"
    repo_url = f"https://github.com/{repo}"
    ev = index_data["electron_version"]

    letter_rows = "\n".join(
        f"| {letter.upper()} | {index_data['files'][f'{letter}.jsonl']:,} | [{letter.upper()}](data/{letter}.md) |"
        for letter in LETTERS
        if f"{letter}.jsonl" in index_data["files"]
    )

    readme = f"""# {name}

已在鸿蒙 PC 的 Electron {ev} 上验证通过的 npm 包列表。这些包可直接在该 Electron 环境中使用，无需额外适配。

## 📊 统计摘要

> 更新日期：{index_data['last_updated']} | 兼容包总数：**{index_data['total']:,}**

## 🔍 在线查询

👉 **[点击打开搜索页]({site_url})**（支持按包名搜索 + 按字母浏览）

**按字母浏览**（GitHub Markdown 表格）：

| 字母 | 包数量 | 浏览表格 |
|------|--------|---------|
{letter_rows}

> 💡 在表格页面中可使用浏览器 `Ctrl+F` 搜索包名。

## ⚠️ 免责声明

兼容性验证结果由 AI 自动验证生成，可能存在部分与实际不符的情况。如发现错误，欢迎 [提交 Issue]({repo_url}/issues) 反馈。

## 📁 项目结构

```
├── index.html               # GitHub Pages 搜索页
├── gen_packages.py          # 数据生成脚本
└── data/
    ├── index.json           # 索引（总条目数、各文件条目数）
    ├── *.md                 # Markdown 表格（GitHub 在线浏览）
    └── *.jsonl              # JSONL 格式（命令行 grep/jq 查询）
```

## 📖 使用方式

### 命令行查询

```bash
git clone {repo_url}.git
cd {name}

# 模糊搜索
grep "lodash" data/l.jsonl

# 精确搜索（scoped 包按 @ 后的首字母分组）
jq -c 'select(.name == "lodash")' data/l.jsonl
jq -c 'select(.name == "@electron/remote")' data/e.jsonl
```

### 程序中引用

```javascript
const fs = require('fs');
const packages = fs.readFileSync('data/l.jsonl', 'utf8')
  .trim().split('\\n').map(JSON.parse);
```

## 🔧 数据更新

```bash
# 1. 准备新的分析结果 xlsx（sheet「{SHEET_NAME}」，含「{COL_STATUS}」「{COL_NAME}」「{COL_VERSION}」列）
# 2. 运行生成脚本（只收录状态为「{PASS_STATUS}」的包）
python3 gen_packages.py <xlsx文件路径> --repo {repo} --electron {ev}
# 3. 提交
git add -A && git commit -m "update: 更新兼容包列表"
git push
```
"""
    with open(README_FILE, "w", encoding="utf-8") as f:
        f.write(readme)


def main():
    parser = argparse.ArgumentParser(description="从 xlsx 生成 Electron 兼容包数据")
    parser.add_argument("xlsx", help="分析结果 xlsx 文件路径")
    parser.add_argument("--repo", default="ohos-electron/compatible-packages", help="GitHub 仓库 owner/name")
    parser.add_argument("--electron", default="34", help="验证所用的 Electron 主版本")
    args = parser.parse_args()

    if not os.path.isfile(args.xlsx):
        sys.exit(f"文件不存在：{args.xlsx}")
    if "/" not in args.repo:
        sys.exit("--repo 格式应为 owner/name")

    print(f"Reading {args.xlsx}...")
    packages = read_packages(args.xlsx)
    index_data = write_data(packages, args.electron, args.repo)
    write_readme(index_data)
    print(f"Done: {index_data['total']} packages, {len(index_data['files'])} files")


if __name__ == "__main__":
    main()
